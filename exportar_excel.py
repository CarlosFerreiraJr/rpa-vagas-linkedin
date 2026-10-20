
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

from extrair_vagas import extrair_vagas


def exportar_excel():
    vagas = extrair_vagas()

    pasta = Path(__file__).resolve().parent
    arquivo = pasta / "vagas_linkedin.xlsx"

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Vagas LinkedIn"

    # Cabeçalho da planilha.
    cabecalho = [
            "ID",
            "Cargo",
            "Empresa",
            "Modalidade",
            "Localização",
            "Descrição",
            "Link da Vaga"
    ]

    planilha.append(cabecalho)

    # Adiciona as vagas coletadas.
    for vaga in vagas.values():
        planilha.append([
            vaga["id"],
            vaga["cargo"],
            vaga["empresa"],
            vaga["modalidade"],
            vaga["localizacao"],
            vaga["descricao"],
            vaga["url"]
        ])

    # Formatação do cabeçalho.
    for celula in planilha[1]:
        celula.font = Font(color="FFFFFF", bold=True)
        celula.fill = PatternFill(
            start_color="17365D",
            end_color="17365D",
            fill_type="solid"
        )
        celula.alignment = Alignment(horizontal="center")

    # Largura das colunas.
    planilha.column_dimensions["A"].width = 18
    planilha.column_dimensions["B"].width = 65
    planilha.column_dimensions["C"].width = 35
    planilha.column_dimensions["D"].width = 20
    planilha.column_dimensions["E"].width = 35
    planilha.column_dimensions["F"].width = 90
    planilha.column_dimensions["G"].width = 55

    # Links clicáveis.
    for linha in range(2, planilha.max_row + 1):
        celula = planilha.cell(row=linha, column=7)
        celula.hyperlink = celula.value
        celula.style = "Hyperlink"

    # Filtros e congelamento do cabeçalho.
    planilha.auto_filter.ref = planilha.dimensions
    planilha.freeze_panes = "A2"

    workbook.save(arquivo)

    print(f"\nPlanilha gerada: {arquivo}")
    print(f"Total de vagas exportadas: {len(vagas)}")


if __name__ == "__main__":
    exportar_excel()
